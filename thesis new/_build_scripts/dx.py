"""Small helper library on top of python-docx for the MCS thesis format (GRP guideline section 8.2)."""
import re
import sys

sys.path.insert(0, "/tmp/claude-1000/-home-krishna-Desktop-Thesis-Hairfall-Thesis/3a0d67b6-c14e-4ab2-95b7-3de697bbdd86/scratchpad/pylibs")
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor

LABELS = {}   # cross-reference key -> "Table 4.7" / "Fig. 4.2" (filled by the previous build pass)
FONT = "Times New Roman"
MONO = "Courier New"
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
C, J, L, R = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT


def set_font(run, size=None, bold=None, italic=None, name=FONT):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def style_font(style, size, bold=False, italic=False, name=FONT):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


class Thesis:
    def __init__(self):
        self.doc = Document()
        self.counters = {}          # ("Table", "4") -> n
        self.captions = {"Table": [], "Fig.": []}   # (label, title)
        self.toc = []               # (level, text, key)
        self.appendices = []
        self.chapter = "0"
        self.labels = {}
        self._next_key = None
        self._setup_styles()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(3), Cm(2.5), Cm(2.5), Cm(2.5)
        sec.footer_distance = Cm(1.2)

    # ------------------------------------------------------------------ styles
    def _setup_styles(self):
        st = self.doc.styles
        n = st["Normal"]
        style_font(n, 12)
        pf = n.paragraph_format
        pf.line_spacing = 1.5
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        pf.alignment = J
        pf.widow_control = True
        for name, size, ital in (("Heading 1", 14, False), ("Heading 2", 12, False), ("Heading 3", 12, False), ("Heading 4", 12, True)):
            s = st[name]
            style_font(s, size, bold=True, italic=ital)
            p = s.paragraph_format
            p.alignment = C if name == "Heading 1" else L
            p.space_before = Pt(12 if name != "Heading 1" else 0)
            p.space_after = Pt(6)
            p.keep_with_next = True
            p.line_spacing = 1.5
        cap = st.add_style("Caption Left", 1)
        style_font(cap, 11, bold=False)
        cap.paragraph_format.alignment = L
        cap.paragraph_format.line_spacing = 1.15
        cap.paragraph_format.space_after = Pt(8)
        cap.paragraph_format.space_before = Pt(4)

    # ------------------------------------------------------------------ basic blocks
    def p(self, text="", align=None, bold=False, italic=False, size=None, space_after=None, keep=False, indent=None, line=None):
        para = self.doc.add_paragraph()
        if align is not None:
            para.alignment = align
        if space_after is not None:
            para.paragraph_format.space_after = Pt(space_after)
        if keep:
            para.paragraph_format.keep_with_next = True
        if indent is not None:
            para.paragraph_format.first_line_indent = Cm(indent)
        if line is not None:
            para.paragraph_format.line_spacing = line
        self._runs(para, text, size=size, bold=bold, italic=italic)
        return para

    def nk(self, key):
        self._next_key = key

    def _take_key(self, num):
        if self._next_key:
            self.labels[self._next_key] = num
            self._next_key = None

    def _runs(self, para, text, size=None, bold=False, italic=False, name=FONT):
        """Inline markup: **bold**, *italic*, ^sup^, ~sub~ (no nesting). {tab:key} / {fig:key} are cross-references."""
        text = re.sub(r"\{(?:tab|fig):(\w+)\}", lambda m: LABELS.get(m.group(1), "?"), text)
        for tok in re.split(r"(\*\*.+?\*\*|\*.+?\*|\^.+?\^|~.+?~)", text):
            if not tok:
                continue
            b, i, sup, sub, t = bold, italic, False, False, tok
            if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
                b, t = True, tok[2:-2]
            elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
                i, t = True, tok[1:-1]
            elif tok.startswith("^") and tok.endswith("^") and len(tok) > 2:
                sup, t = True, tok[1:-1]
            elif tok.startswith("~") and tok.endswith("~") and len(tok) > 2:
                sub, t = True, tok[1:-1]
            r = para.add_run(t)
            set_font(r, size=size, bold=b, italic=i, name=name)
            if sup:
                r.font.superscript = True
            if sub:
                r.font.subscript = True

    def blank(self, n=1, size=14):
        for _ in range(n):
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            r = para.add_run(""); set_font(r, size=size)

    def bullets(self, items, numbered=False, space=3):
        for k, t in enumerate(items, 1):
            para = self.doc.add_paragraph()
            para.paragraph_format.left_indent = Cm(1.0)
            para.paragraph_format.first_line_indent = Cm(-0.7)
            para.paragraph_format.space_after = Pt(space)
            self._runs(para, (f"{k}. " if numbered else "•  ") + t)

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------ headings and TOC registration
    def _reg(self, level, text, key=None):
        self.toc.append((level, text, key or text))

    def chapter_heading(self, num, title):
        self.chapter = str(num)
        h = self.doc.add_paragraph(style="Heading 1")
        h.paragraph_format.page_break_before = True
        h.paragraph_format.space_after = Pt(14)
        r = h.add_run(f"CHAPTER {num}"); set_font(r, 14, bold=True)
        r.add_break()
        r2 = h.add_run(title.upper()); set_font(r2, 14, bold=True)
        self._reg(0, f"CHAPTER {num}  {title.upper()}", f"CHAPTER {num}")

    def appendix_heading(self, letter, title):
        self.chapter = letter
        h = self.doc.add_paragraph(style="Heading 1")
        h.paragraph_format.page_break_before = True
        h.paragraph_format.space_after = Pt(14)
        r = h.add_run(f"APPENDIX {letter}"); set_font(r, 14, bold=True)
        r.add_break()
        r2 = h.add_run(title.upper()); set_font(r2, 14, bold=True)
        self._reg(0, f"Appendix {letter}: {title}", f"APPENDIX {letter}")
        self.appendices.append((f"Appendix {letter}:", title))

    def unnumbered_heading(self, text, toc=True, page_break=True, level=0):
        h = self.doc.add_paragraph(style="Heading 1")
        h.paragraph_format.page_break_before = page_break
        h.paragraph_format.space_after = Pt(14)
        r = h.add_run(text); set_font(r, 14, bold=True)
        if toc:
            self._reg(level, text, text)

    def h2(self, text):
        self.doc.add_paragraph(style="Heading 2").add_run(text)
        self._fix_last_heading()
        self._reg(1, text, text)

    def h3(self, text):
        self.doc.add_paragraph(style="Heading 3").add_run(text)
        self._fix_last_heading()
        self._reg(2, text, text)

    def h4(self, text):
        self.doc.add_paragraph(style="Heading 4").add_run(text)
        self._fix_last_heading()

    def _fix_last_heading(self):
        para = self.doc.paragraphs[-1]
        for r in para.runs:
            set_font(r, 12, bold=True, italic=(para.style.name == "Heading 4"))

    # ------------------------------------------------------------------ captions
    def _number(self, label):
        key = (label, self.chapter)
        self.counters[key] = self.counters.get(key, 0) + 1
        return f"{label} {self.chapter}.{self.counters[key]}"

    def table(self, caption, header, rows, widths, size=10, align=None, bold_last=False, note=None, header_fill="D9D9D9", label="Table"):
        num = self._number(label)
        self._take_key(num)
        self.captions["Table"].append((num, caption))
        cp = self.doc.add_paragraph(style="Caption Left")
        cp.paragraph_format.keep_with_next = True
        self._runs(cp, f"{num}: {caption}", size=11)
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        self._borders(t)
        align = align or ["l"] + ["c"] * (len(header) - 1)
        for ri, row in enumerate([header] + rows):
            for ci, val in enumerate(row):
                cell = t.cell(ri, ci)
                cell.width = Cm(widths[ci])
                para = cell.paragraphs[0]
                para.paragraph_format.line_spacing = 1.0
                para.paragraph_format.space_after = Pt(1)
                para.paragraph_format.space_before = Pt(1)
                para.alignment = {"l": L, "c": C, "r": R}[align[ci]]
                self._runs(para, str(val), size=size, bold=(ri == 0) or (bold_last and ri == len(rows)))
                if ri == 0:
                    tcpr = cell._element.get_or_add_tcPr()
                    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), header_fill)
                    tcpr.append(shd)
            trpr = t.rows[ri]._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trpr.append(cs)
            if ri == 0:
                th = OxmlElement("w:tblHeader"); trpr.append(th)
        if len(rows) <= 16:
            for ri in range(len(rows)):
                for cell in t.rows[ri].cells:
                    for para in cell.paragraphs:
                        para.paragraph_format.keep_with_next = True
        if note:
            np_ = self.doc.add_paragraph(); np_.paragraph_format.line_spacing = 1.15; np_.paragraph_format.space_before = Pt(3)
            self._runs(np_, note, size=10, italic=False)
        else:
            sp = self.doc.add_paragraph(); sp.paragraph_format.space_after = Pt(4); sp.paragraph_format.line_spacing = 0.8
        return num

    def _borders(self, t):
        tblPr = t._tbl.tblPr
        b = parse_xml(f'<w:tblBorders {nsdecls("w")}>' + "".join(
            f'<w:{e} w:val="single" w:sz="4" w:space="0" w:color="000000"/>' for e in ("top", "left", "bottom", "right", "insideH", "insideV")) + "</w:tblBorders>")
        mar = parse_xml(f'<w:tblCellMar {nsdecls("w")}><w:left w:w="70" w:type="dxa"/><w:right w:w="70" w:type="dxa"/></w:tblCellMar>')
        layout, look = tblPr.find(qn("w:tblLayout")), tblPr.find(qn("w:tblLook"))
        if layout is not None:
            layout.addprevious(b)          # schema order: tblBorders < tblLayout < tblCellMar < tblLook
            layout.addnext(mar)
        elif look is not None:
            look.addprevious(b); look.addprevious(mar)
        else:
            tblPr.append(b); tblPr.append(mar)

    def figure(self, path, caption, width_cm=14.0, label="Fig."):
        num = self._number(label)
        self._take_key(num)
        self.captions["Fig."].append((num, caption))
        para = self.doc.add_paragraph()
        para.alignment = C
        para.paragraph_format.keep_with_next = True
        para.paragraph_format.line_spacing = 1.0
        para.paragraph_format.space_after = Pt(2)
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cp = self.doc.add_paragraph(style="Caption Left")
        self._runs(cp, f"{num}: {caption}", size=11)
        return num

    def figure_row(self, paths, caption, width_cm=4.9, label="Fig.", subcaps=None):
        """Several images side by side sharing one caption."""
        num = self._number(label)
        self._take_key(num)
        self.captions["Fig."].append((num, caption))
        t = self.doc.add_table(rows=2 if subcaps else 1, cols=len(paths))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, pth in enumerate(paths):
            para = t.cell(0, i).paragraphs[0]
            para.alignment = C; para.paragraph_format.line_spacing = 1.0; para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(pth), width=Cm(width_cm))
            if subcaps:
                sp = t.cell(1, i).paragraphs[0]
                sp.alignment = C; sp.paragraph_format.line_spacing = 1.0; sp.paragraph_format.space_after = Pt(0)
                self._runs(sp, subcaps[i], size=10)
        cp = self.doc.add_paragraph(style="Caption Left")
        cp.paragraph_format.space_before = Pt(6)
        self._runs(cp, f"{num}: {caption}", size=11)
        return num

    def figure_grid(self, paths, caption, cols=2, width_cm=7.6, subcaps=None, label="Fig."):
        """Images in a grid (cols per row) sharing one caption."""
        num = self._number(label)
        self._take_key(num)
        self.captions["Fig."].append((num, caption))
        nrows = -(-len(paths) // cols)
        t = self.doc.add_table(rows=nrows * (2 if subcaps else 1), cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, pth in enumerate(paths):
            r, c = divmod(i, cols)
            ri = r * (2 if subcaps else 1)
            para = t.cell(ri, c).paragraphs[0]
            para.alignment = C; para.paragraph_format.line_spacing = 1.0; para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.keep_with_next = True
            para.add_run().add_picture(str(pth), width=Cm(width_cm))
            if subcaps:
                sp = t.cell(ri + 1, c).paragraphs[0]
                sp.alignment = C; sp.paragraph_format.line_spacing = 1.0; sp.paragraph_format.space_after = Pt(4)
                sp.paragraph_format.keep_with_next = (r < nrows - 1)
                self._runs(sp, subcaps[i], size=10)
        cp = self.doc.add_paragraph(style="Caption Left")
        cp.paragraph_format.space_before = Pt(6)
        self._runs(cp, f"{num}: {caption}", size=11)
        return num

    # ------------------------------------------------------------------ equations (OMML)
    def equation(self, omml_inner, number):
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.alignment = L
        pf.line_spacing = 1.15
        pf.space_before = Pt(4); pf.space_after = Pt(8)
        pf.tab_stops.add_tab_stop(Cm(7.75), WD_TAB_ALIGNMENT.CENTER)
        pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT)
        r = para.add_run("\t"); set_font(r, 12)
        el = parse_xml(f'<m:oMath xmlns:m="{NS_M}" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">{omml_inner}</m:oMath>')
        para._p.append(el)
        r = para.add_run(f"\t({number})"); set_font(r, 12)

    # ------------------------------------------------------------------ code
    def code(self, text, size=7.5):
        for line in text.rstrip("\n").split("\n"):
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.line_spacing = 1.0; pf.space_after = Pt(0); pf.alignment = L
            pf.left_indent = Cm(0.4); pf.first_line_indent = Cm(-0.0)
            r = para.add_run(line.replace("\t", "    ") if line else " ")
            set_font(r, size, name=MONO)
        self.blank(1, 6)

    # ------------------------------------------------------------------ sections and page numbers
    def new_section(self, fmt=None, start=None, footer=True, first_page_blank=False):
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._pgnum(sec, fmt, start)
        sec.footer.is_linked_to_previous = False
        self._clear_footer(sec.footer)
        sec.different_first_page_header_footer = first_page_blank
        if first_page_blank:
            sec.first_page_footer.is_linked_to_previous = False
            self._clear_footer(sec.first_page_footer)
        if footer:
            self._page_field(sec.footer)
        return sec

    def _pgnum(self, sec, fmt, start):
        sectPr = sec._sectPr
        old = sectPr.find(qn("w:pgNumType"))
        if old is not None:
            sectPr.remove(old)
        pg = OxmlElement("w:pgNumType")
        if fmt:
            pg.set(qn("w:fmt"), fmt)
        if start is not None:
            pg.set(qn("w:start"), str(start))
        mar = sectPr.find(qn("w:pgMar"))
        if mar is not None:
            mar.addnext(pg)                # schema order: pgSz, pgMar, ..., pgNumType, cols
        else:
            sectPr.append(pg)

    def _clear_footer(self, footer):
        for para in footer.paragraphs:
            for r in list(para.runs):
                r._element.getparent().remove(r._element)

    def _page_field(self, footer):
        para = footer.paragraphs[0]
        para.alignment = C
        for kind, txt in (("begin", None), (None, " PAGE "), ("end", None)):
            r = para.add_run(); set_font(r, 12)
            if kind:
                fc = OxmlElement("w:fldChar"); fc.set(qn("w:fldCharType"), kind); r._element.append(fc)
            else:
                it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = txt; r._element.append(it)

    def first_section_setup(self, footer=False):
        sec = self.doc.sections[0]
        self._clear_footer(sec.footer)

    # ------------------------------------------------------------------ static contents pages
    def toc_page(self, title, entries, pages, size=12, level_indent=(0, 0.0, 0.8, 1.6)):
        self.unnumbered_heading(title, toc=False)
        for lvl, text, key in entries:
            self._toc_line(lvl, text, pages.get(("toc", key), ""), size, level_indent)

    def _toc_line(self, lvl, text, pg, size=12, level_indent=(0, 0.0, 0.8, 1.6)):
        para = self.doc.add_paragraph()
        pf = para.paragraph_format
        pf.alignment = L
        pf.line_spacing = 1.0
        pf.space_after = Pt(3 if lvl != 0 else 2)
        pf.space_before = Pt(8 if lvl == 0 else 0)
        pf.keep_with_next = lvl == 0
        ind = {0: 0.0, 0.5: 0.0, 1: 0.5, 2: 1.3}[lvl]
        pf.left_indent = Cm(ind + 1.0)
        pf.first_line_indent = Cm(-1.0)
        pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        bold = lvl in (0, 0.5)
        r = para.add_run(text + ("\t" + str(pg) if lvl != 0 else "\t" + str(pg) if pg else "\t"))
        set_font(r, size, bold=bold)

    def list_page(self, title, items, pages, kind):
        self.unnumbered_heading(title, toc=False)
        for num, cap in items:
            para = self.doc.add_paragraph()
            pf = para.paragraph_format
            pf.line_spacing = 1.15; pf.space_after = Pt(5); pf.alignment = L
            pf.left_indent = Cm(2.4); pf.first_line_indent = Cm(-2.4)
            pf.tab_stops.add_tab_stop(Cm(2.4), WD_TAB_ALIGNMENT.LEFT)
            pf.tab_stops.add_tab_stop(Cm(15.5), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            r = para.add_run(f"{num}\t{cap}\t{pages.get((kind, num), '')}"); set_font(r, 12)

    def save(self, path):
        self.doc.save(path)


# ---------------------------------------------------------------------- OMML mini builder
def _rpr(plain):
    return '<m:rPr><m:sty m:val="p"/></m:rPr>' if plain else ""


def mr(text, plain=False):
    t = (text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    return (f'<m:r>{_rpr(plain)}<w:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/></w:rPr>'
            f'<m:t xml:space="preserve">{t}</m:t></m:r>')


def msub(base, sub):
    return f"<m:sSub><m:e>{base}</m:e><m:sub>{sub}</m:sub></m:sSub>"


def msup(base, sup):
    return f"<m:sSup><m:e>{base}</m:e><m:sup>{sup}</m:sup></m:sSup>"


def msubsup(base, sub, sup):
    return f"<m:sSubSup><m:e>{base}</m:e><m:sub>{sub}</m:sub><m:sup>{sup}</m:sup></m:sSubSup>"


def mfrac(num, den):
    return f"<m:f><m:num>{num}</m:num><m:den>{den}</m:den></m:f>"


def msqrt(x):
    return f'<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{x}</m:e></m:rad>'


def mnary(sym, lo, hi, body):
    hide = ("" if lo else '<m:subHide m:val="1"/>') + ("" if hi else '<m:supHide m:val="1"/>')
    return (f'<m:nary><m:naryPr><m:chr m:val="{sym}"/><m:limLoc m:val="undOvr"/>{hide}</m:naryPr>'
            f"<m:sub>{lo}</m:sub><m:sup>{hi}</m:sup><m:e>{body}</m:e></m:nary>")


def mdelim(inner, beg="(", end=")"):
    return f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr><m:e>{inner}</m:e></m:d>'
